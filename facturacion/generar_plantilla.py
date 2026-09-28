# -*- coding: utf-8 -*-
"""
Plantilla de solicitud de información para FACTURACIÓN POR LÍNEA DE LABOR (ASM).
Todo va PRELLENADO con la propuesta; el cliente solo confirma o corrige (celdas amarillas)
y escoge de listas: así el archivo devuelto se carga solo a la app, sin redigitar.
Los nombres de hojas y encabezados son FIJOS: el importador los lee tal cual.
"""
import csv, io, os, json
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

BASE = os.path.dirname(os.path.abspath(__file__))
SALIDA = os.path.join(BASE, 'ASM_solicitud_tarifas_facturacion.xlsx')

VERDE = '1A6B3A'; VERDE_CLARO = 'E6F0E9'; AMARILLO = 'FFF2CC'; GRIS = 'F2F2F2'; ROJO = 'FCE4E4'; NARANJA = 'FDEBD3'
fino = Side(style='thin', color='BFBFBF')
BORDE = Border(left=fino, right=fino, top=fino, bottom=fino)

RS = ['AGROMORALES', 'CEBALLOS Y LOZANO']
CLIENTES = ['Ingenio San Carlos', 'Ingenio Riopaila', 'Riopaila Agrícola', 'Ingenio Mayagüez', 'Ingenio Risaralda',
            'Ingenio Pichichí', 'Proveedor', 'Ingenio Carmelita', 'Trapiche Lucerna', 'Ingenio Providencia', 'Otro (ver observación)']
ING = {'san_carlos': 'Ingenio San Carlos', 'riopaila': 'Ingenio Riopaila', 'mayaguez': 'Ingenio Mayagüez',
       'risaralda': 'Ingenio Risaralda', 'pichichi': 'Ingenio Pichichí', 'ingenio_carmelita': 'Ingenio Carmelita',
       'trapiche_lucerna': 'Trapiche Lucerna', 'proveedor': 'Proveedor', '': '(sin ingenio en la app)'}
UNIDADES = ['Hectárea', 'Hectómetro', 'Hora', 'Jornal']
LABORES_APP = ['DESPEJE', 'DESPEJE 0 X 0', 'REENCALLE', 'REENCALLE V', 'TRIPLE', 'TRIPLE 4X1 (nueva)', 'SUBSUELO',
               'SUBSUELO 4X1 (nueva)', 'FERTILIZACION PLANTILLA (nueva)', 'FERTILIZACION 2X1 (nueva)',
               'FERTILIZACION 4X1 (nueva)', 'CULTIVO', 'CULTIVO PLANTILLA (nueva)', 'CULTIVO 2X1 (nueva)', 'ACEQUIAS',
               'DEVOLUCIÓN DE BASURA', 'OFICIOS VARIOS (hora)', 'OFICIOS VARIOS JORNAL (nueva)', 'Otra (ver observación)']
CONFIRMA = ['Sí, correcto', 'Corregir (ver columnas)', 'Eliminar esta fila']
SINO = ['Sí', 'No']

# ── La tabla que mandó el cliente (imagen del 28-sep-2026) ───────────────────
# (razón social, cliente, labor como la escriben, labor en la app, unidad, precio)
TARIFAS = [
    ('AGROMORALES', 'Ingenio San Carlos', 'Despeje', 'DESPEJE', 'Hectárea', 85786),
    ('AGROMORALES', 'Ingenio San Carlos', 'Reencalle', 'REENCALLE', 'Hectárea', 85786),
    ('AGROMORALES', 'Ingenio Riopaila', 'Despeje', 'DESPEJE', 'Hectárea', 107203),
    ('AGROMORALES', 'Ingenio Riopaila', 'Reencalle', 'REENCALLE', 'Hectárea', 97364),
    ('AGROMORALES', 'Ingenio Riopaila', 'Reencalle caña cruda', 'REENCALLE V', 'Hectárea', 181728),
    ('AGROMORALES', 'Ingenio Riopaila', 'Triple', 'TRIPLE', 'Hectárea', 187554),
    ('AGROMORALES', 'Ingenio Riopaila', 'Cultivo aporque', 'CULTIVO', 'Hectárea', 139678),
    ('AGROMORALES', 'Ingenio Riopaila', 'Of, varios Jornal', 'OFICIOS VARIOS JORNAL (nueva)', 'Jornal', 640000),
    ('AGROMORALES', 'Riopaila Agrícola', 'Reencalle', 'REENCALLE', 'Hectárea', 111000),
    ('AGROMORALES', 'Riopaila Agrícola', 'Abono Plantilla', 'FERTILIZACION PLANTILLA (nueva)', 'Hectárea', 167214),
    ('AGROMORALES', 'Riopaila Agrícola', 'Abono 2x1', 'FERTILIZACION 2X1 (nueva)', 'Hectárea', 159000),
    ('AGROMORALES', 'Riopaila Agrícola', 'Abono 4x1', 'FERTILIZACION 4X1 (nueva)', 'Hectárea', 165563),
    ('AGROMORALES', 'Riopaila Agrícola', 'Acequias', 'ACEQUIAS', None, 70000),
    ('AGROMORALES', 'Riopaila Agrícola', 'Of, varios Hora', 'OFICIOS VARIOS (hora)', 'Hora', 50000),
    ('AGROMORALES', 'Riopaila Agrícola', 'Subsuelo', 'SUBSUELO', 'Hectárea', 260000),
    ('AGROMORALES', 'Riopaila Agrícola', 'Triple', 'TRIPLE', 'Hectárea', 219764),
    ('AGROMORALES', 'Riopaila Agrícola', 'Cultivo aporque', 'CULTIVO', 'Hectárea', 106900),
    ('AGROMORALES', 'Ingenio Mayagüez', 'Despeje', 'DESPEJE', 'Hectárea', 85786),
    ('AGROMORALES', 'Ingenio Mayagüez', 'Reencalle', 'REENCALLE', 'Hectárea', 85786),
    ('AGROMORALES', 'Ingenio Risaralda', 'Despeje', 'DESPEJE', 'Hectárea', 105574),
    ('AGROMORALES', 'Ingenio Risaralda', 'Reencalle', 'REENCALLE', 'Hectárea', 97859),
    ('CEBALLOS Y LOZANO', 'Ingenio Pichichí', 'Despeje', 'DESPEJE', 'Hectárea', 82504),
    ('CEBALLOS Y LOZANO', 'Ingenio Pichichí', 'Reencalle', 'REENCALLE', 'Hectárea', 82504),
    ('CEBALLOS Y LOZANO', 'Ingenio Pichichí', 'Subsuelo', 'SUBSUELO', 'Hectárea', 208489),
    ('CEBALLOS Y LOZANO', 'Ingenio Pichichí', 'Triple', 'TRIPLE', 'Hectárea', 169003),
    ('CEBALLOS Y LOZANO', 'Ingenio Riopaila', 'Despeje', 'DESPEJE', 'Hectárea', 108028),
    ('CEBALLOS Y LOZANO', 'Proveedor', 'Reencalle', 'REENCALLE', 'Hectárea', 105000),
    ('CEBALLOS Y LOZANO', 'Proveedor', 'Reencalle caña verde', 'REENCALLE V', 'Hectárea', 190000),
    ('CEBALLOS Y LOZANO', 'Proveedor', 'Subsuelo', 'SUBSUELO', 'Hectárea', 225000),
    ('CEBALLOS Y LOZANO', 'Proveedor', 'Triple', 'TRIPLE', 'Hectárea', 225000),
    ('CEBALLOS Y LOZANO', 'Proveedor', 'Subsuelo 4x1', 'SUBSUELO 4X1 (nueva)', 'Hectárea', 240000),
    ('CEBALLOS Y LOZANO', 'Proveedor', 'Triple 4x1', 'TRIPLE 4X1 (nueva)', 'Hectárea', 240000),
    ('CEBALLOS Y LOZANO', 'Proveedor', 'Abono Plantilla', 'FERTILIZACION PLANTILLA (nueva)', 'Hectárea', 170000),
    ('CEBALLOS Y LOZANO', 'Proveedor', 'Abono 2x1', 'FERTILIZACION 2X1 (nueva)', 'Hectárea', 160000),
    ('CEBALLOS Y LOZANO', 'Proveedor', 'Cultivo aporque plantilla', 'CULTIVO PLANTILLA (nueva)', 'Hectárea', 130000),
    ('CEBALLOS Y LOZANO', 'Proveedor', 'Cultivo aporque 2x1', 'CULTIVO 2X1 (nueva)', 'Hectárea', 120000),
    ('CEBALLOS Y LOZANO', 'Proveedor', 'Acequias', 'ACEQUIAS', None, 85000),
]

# ¿Qué precio(s) de la tabla le caen a cada labor de la app, según el cliente de la tarifa?
BASE_LABOR = {'DESPEJE': ['DESPEJE'], 'REENCALLE': ['REENCALLE'], 'REENCALLE V': ['REENCALLE V'],
              'TRIPLE': ['TRIPLE', 'TRIPLE 4X1 (nueva)'], 'SUBSUELO': ['SUBSUELO', 'SUBSUELO 4X1 (nueva)'],
              'FERTILIZACION': ['FERTILIZACION PLANTILLA (nueva)', 'FERTILIZACION 2X1 (nueva)', 'FERTILIZACION 4X1 (nueva)'],
              'CULTIVO': ['CULTIVO', 'CULTIVO PLANTILLA (nueva)', 'CULTIVO 2X1 (nueva)'],
              'ACEQUIAS': ['ACEQUIAS'], 'DESPEJE 0 X 0': [], 'DEVOLUCIÓN DE BASURA': []}

def precios(cliente_tarifa, labor_app):
    claves = BASE_LABOR.get(labor_app, [])
    return [t for t in TARIFAS if t[1] == cliente_tarifa and t[3] in claves]

def clasificar(ingenio, tipo, labor):
    """Propuesta de regla → (razón social propuesta, estado, motivo)."""
    if tipo == 'proveedores':
        cands = precios('Proveedor', labor); rs = 'CEBALLOS Y LOZANO'
    elif ingenio == 'pichichi':
        cands = precios('Ingenio Pichichí', labor); rs = 'CEBALLOS Y LOZANO'
    elif ingenio == 'riopaila':
        a = precios('Ingenio Riopaila', labor); b = precios('Riopaila Agrícola', labor)
        rs = 'AGROMORALES'
        if a and b:
            return rs, 'Falta decidir', 'Hay precio en Ingenio Riopaila Y en Riopaila Agrícola: depende de la hacienda (hoja 4)'
        cands = a or b
        if labor == 'DESPEJE':
            return rs, 'Falta decidir', 'Despeje de Riopaila tiene precio en AGROMORALES ($107.203) y en CEBALLOS Y LOZANO ($108.028): ¿cuál?'
    elif ingenio in ('san_carlos', 'mayaguez', 'risaralda'):
        cands = precios(ING[ingenio], labor); rs = 'AGROMORALES'
    else:
        return '', 'Sin precio', 'No hay tabla de precios para este cliente'
    if not cands:
        return rs, 'Sin precio', 'La tabla no trae precio de esta labor para este cliente'
    if len(cands) > 1:
        return rs, 'Falta decidir', 'Hay varias variantes (' + ' / '.join(c[2] for c in cands) + '): ¿cómo se sabe cuál es? (hoja 5)'
    if cands[0][4] is None:
        return rs, 'Falta decidir', 'Falta la unidad de cobro (¿hectárea o hectómetro?)'
    if labor == 'REENCALLE V':
        return rs, 'Con precio', f'Se asume REENCALLE V = «{cands[0][2]}» — confirmar'
    return rs, 'Con precio', ''

def leer(nombre):
    with io.open(os.path.join(BASE, nombre), encoding='utf-8') as f:
        return [r for r in csv.DictReader(l for l in f if l.strip() and not l.startswith('Shell cwd'))]

combos = leer('combos.csv'); haciendas = leer('haciendas.csv')

wb = Workbook()

def encabezado(ws, fila, titulos, anchos):
    for i, (t, w) in enumerate(zip(titulos, anchos), start=1):
        c = ws.cell(row=fila, column=i, value=t)
        c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor=VERDE)
        c.alignment = Alignment(wrap_text=True, vertical='center'); c.border = BORDE
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[fila].height = 42
    ws.freeze_panes = ws.cell(row=fila + 1, column=1)

def titulo(ws, texto, sub):
    ws['A1'] = texto; ws['A1'].font = Font(bold=True, size=14, color=VERDE)
    ws['A2'] = sub; ws['A2'].font = Font(italic=True, size=10, color='595959')
    ws['A2'].alignment = Alignment(wrap_text=False)

def celda(ws, f, c, v, llenar=False, fmt=None, color=None):
    x = ws.cell(row=f, column=c, value=v); x.border = BORDE
    x.alignment = Alignment(vertical='top', wrap_text=isinstance(v, str) and len(v) > 30)
    if llenar: x.fill = PatternFill('solid', fgColor=AMARILLO)
    elif color: x.fill = PatternFill('solid', fgColor=color)
    if fmt: x.number_format = fmt
    return x

def lista(ws, rango, hoja_lista, n, ayuda=None):
    dv = DataValidation(type='list', formula1=f"=Listas!${hoja_lista}$2:${hoja_lista}${n + 1}", allow_blank=True)
    dv.error = 'Escoja un valor de la lista.'; dv.errorTitle = 'Valor no válido'
    if ayuda: dv.prompt = ayuda; dv.showInputMessage = True
    ws.add_data_validation(dv); dv.add(rango)

PESOS = '"$"#,##0'; HA = '#,##0.00'

# ── Léame ──────────────────────────────────────────────────────────────────
ws = wb.active; ws.title = 'Léame'
ws.column_dimensions['A'].width = 110
lineas = [
    ('AgroServicios Morales · Información para facturar cada labor realizada', 'h1'),
    ('Para qué es: que la app calcule sola el valor a facturar de CADA labor terminada (hectáreas × tarifa), con la razón social y el cliente correctos, y arme la pre-factura por quincena.', ''),
    ('', ''),
    ('CÓMO LLENARLO', 'h2'),
    ('• Todo viene PRELLENADO con nuestra propuesta. Solo revise y corrija lo que no esté bien.', ''),
    ('• Las celdas AMARILLAS son las que se llenan o se confirman. Las blancas y grises no se tocan.', ''),
    ('• Donde hay lista desplegable, escoja de la lista (no escriba a mano): así el archivo se carga solo a la app.', ''),
    ('• Si algo no cabe en la lista, escoja «Otro» y explíquelo en la columna Observación.', ''),
    ('• No cambie los nombres de las hojas ni de las columnas.', ''),
    ('', ''),
    ('LAS HOJAS (en orden de importancia)', 'h2'),
    ('1 Tarifas — la tabla de precios que nos pasaron, ya ordenada. Confirmar cada fila, poner la fecha desde la que rige y si incluye IVA.', ''),
    ('2 Faltan precios — labores que SÍ se han hecho (registradas en la app desde mayo) y que la tabla no alcanza a precisar. Ordenadas de más a menos hectáreas.', ''),
    ('3 Razón social — con qué razón social se factura a cada cliente (AGROMORALES o CEBALLOS Y LOZANO).', ''),
    ('4 Haciendas — solo si la razón social o el cliente cambian según la hacienda (sobre todo Riopaila vs Riopaila Agrícola).', ''),
    ('5 Labores y variantes — cómo se sabe si un abono es plantilla, 2x1 o 4x1, y otras dudas de labores.', ''),
    ('6 Preguntas — cómo se factura (cada cuánto, qué lleva la factura, IVA…).', ''),
    ('', ''),
    ('QUÉ PASA DESPUÉS', 'h2'),
    ('Nos devuelven este mismo archivo y se carga directo: cada labor queda con su valor en la app, las que no tengan precio salen marcadas «sin tarifa» (nunca en $0 escondido), y cambiar un precio más adelante NO cambia lo ya facturado.', ''),
]
for i, (t, tipo) in enumerate(lineas, start=1):
    c = ws.cell(row=i, column=1, value=t); c.alignment = Alignment(wrap_text=True, vertical='top')
    if tipo == 'h1': c.font = Font(bold=True, size=15, color=VERDE)
    elif tipo == 'h2': c.font = Font(bold=True, size=11, color=VERDE)
ley = len(lineas) + 2
ws.cell(row=ley, column=1, value='Celda para llenar o confirmar').fill = PatternFill('solid', fgColor=AMARILLO)
ws.cell(row=ley + 1, column=1, value='Dato que sale de la app (no tocar)').fill = PatternFill('solid', fgColor=GRIS)

# ── 1 Tarifas ──────────────────────────────────────────────────────────────
ws = wb.create_sheet('1 Tarifas')
titulo(ws, '1 · Tarifas', 'La tabla que nos pasaron, ordenada. Confirme cada fila (col. I), la fecha desde la que rige (H) y si incluye IVA (G). Agregue abajo los precios que falten.')
cols = ['Razón social que factura', 'Cliente', 'Labor (como la escriben ustedes)', 'Labor en la app', 'Unidad de cobro',
        'Precio unitario (COP)', '¿El precio incluye IVA?', 'Vigente desde (dd/mm/aaaa)', '¿Está correcto?', 'Observación']
encabezado(ws, 4, cols, [22, 22, 28, 34, 16, 16, 14, 16, 22, 40])
f = 5
for rs, cli, lab, app, uni, pr in TARIFAS:
    celda(ws, f, 1, rs, color=GRIS); celda(ws, f, 2, cli, color=GRIS); celda(ws, f, 3, lab, color=GRIS)
    celda(ws, f, 4, app, llenar=True); celda(ws, f, 5, uni, llenar=True); celda(ws, f, 6, pr, llenar=False, fmt=PESOS, color=GRIS)
    celda(ws, f, 7, None, llenar=True); celda(ws, f, 8, None, llenar=True, fmt='dd/mm/yyyy'); celda(ws, f, 9, None, llenar=True)
    obs = ''
    if uni is None: obs = '¿Se cobra por hectárea o por hectómetro? (en la app hay registros de las dos formas)'
    elif app == 'REENCALLE V': obs = 'Suponemos que la «REENCALLE V» de la app es este reencalle de caña verde/cruda: confirmar'
    elif '(nueva)' in app: obs = 'Labor nueva en la app, para que el supervisor escoja la variante al asignar'
    celda(ws, f, 10, obs, llenar=True)
    f += 1
fin_tabla = f
for _ in range(25):  # filas libres para lo que falte
    for c in range(1, 11):
        celda(ws, f, c, None, llenar=True, fmt=PESOS if c == 6 else ('dd/mm/yyyy' if c == 8 else None))
    f += 1
ultima = f - 1
lista(ws, f'A5:A{ultima}', 'A', len(RS)); lista(ws, f'B5:B{ultima}', 'B', len(CLIENTES))
lista(ws, f'D5:D{ultima}', 'C', len(LABORES_APP)); lista(ws, f'E5:E{ultima}', 'D', len(UNIDADES))
lista(ws, f'G5:G{ultima}', 'F', len(SINO)); lista(ws, f'I5:I{fin_tabla - 1}', 'E', len(CONFIRMA))
dv = DataValidation(type='decimal', operator='greaterThan', formula1='0', allow_blank=True)
dv.error = 'Escriba el precio en pesos, sin puntos ni signo.'; ws.add_data_validation(dv); dv.add(f'F5:F{ultima}')
dvf = DataValidation(type='date', operator='greaterThan', formula1='DATE(2025,1,1)', allow_blank=True)
dvf.error = 'Escriba una fecha (dd/mm/aaaa).'; ws.add_data_validation(dvf); dvf.add(f'H5:H{ultima}')
ws.cell(row=fin_tabla, column=1).value = None
ws.cell(row=3, column=1, value=f'Filas {5}–{fin_tabla - 1}: tabla recibida el 28-sep-2026 · Filas {fin_tabla}–{ultima}: libres para agregar precios que falten').font = Font(size=9, color='7F7F7F')

# ── 2 Faltan precios ───────────────────────────────────────────────────────
ws = wb.create_sheet('2 Faltan precios')
titulo(ws, '2 · Labores hechas que la tabla no alcanza a precisar',
       'Salen de la app (labores terminadas desde el 16-may-2026). Complete las columnas amarillas. Ordenadas de más a menos cantidad.')
cols = ['Cliente (en la app)', 'Tipo', 'Labor en la app', 'Unidad en la app', 'Labores registradas', 'Cantidad realizada',
        'Desde', 'Hasta', 'Qué falta', 'Razón social que factura', 'Cliente que se factura', 'Labor / variante que aplica',
        'Precio unitario (COP)', 'Unidad de cobro', 'Observación']
encabezado(ws, 4, cols, [22, 12, 22, 10, 11, 12, 11, 11, 52, 20, 20, 30, 15, 14, 36])
resumen = defaultdict(float)
pendientes = []
for r in combos:
    ing, tipo, lab = r['ingenio'], r['cliente'], r['labor']
    rs, estado, motivo = clasificar(ing, tipo, lab)
    cant = float(r['cantidad'] or 0)
    resumen[estado] += cant
    if estado != 'Con precio':
        pendientes.append((estado, cant, r, rs, motivo))
pendientes.sort(key=lambda x: (-x[1]))
f = 5
for estado, cant, r, rs, motivo in pendientes:
    tipo = 'Proveedor' if r['cliente'] == 'proveedores' else ('Ingenio' if r['cliente'] == 'ingenios' else '(vacío)')
    color = ROJO if estado == 'Sin precio' else NARANJA
    celda(ws, f, 1, ING.get(r['ingenio'], r['ingenio']), color=GRIS); celda(ws, f, 2, tipo, color=GRIS)
    celda(ws, f, 3, r['labor'], color=GRIS); celda(ws, f, 4, r['unidad'], color=GRIS)
    celda(ws, f, 5, int(r['lineas']), color=GRIS); celda(ws, f, 6, cant, fmt=HA, color=GRIS)
    celda(ws, f, 7, r['desde'], color=GRIS); celda(ws, f, 8, r['hasta'], color=GRIS)
    celda(ws, f, 9, f'{estado.upper()}: {motivo}', color=color)
    celda(ws, f, 10, rs or None, llenar=True); celda(ws, f, 11, None, llenar=True); celda(ws, f, 12, None, llenar=True)
    celda(ws, f, 13, None, llenar=True, fmt=PESOS); celda(ws, f, 14, None, llenar=True); celda(ws, f, 15, None, llenar=True)
    f += 1
ult = f - 1
lista(ws, f'J5:J{ult}', 'A', len(RS)); lista(ws, f'K5:K{ult}', 'B', len(CLIENTES))
lista(ws, f'L5:L{ult}', 'C', len(LABORES_APP)); lista(ws, f'N5:N{ult}', 'D', len(UNIDADES))
ws.cell(row=3, column=1, value='Rojo = no hay precio en la tabla · Naranja = hay precio pero falta decidir cuál').font = Font(size=9, color='7F7F7F')

# ── 3 Razón social ─────────────────────────────────────────────────────────
ws = wb.create_sheet('3 Razón social')
titulo(ws, '3 · Con qué razón social se factura a cada cliente',
       'Propuesta según la tabla recibida. Si depende de la hacienda, escoja «Depende de la hacienda» y marque las haciendas en la hoja 4.')
cols = ['Cliente (en la app)', 'Tipo', 'Labores registradas', 'Cantidad realizada', 'Propuesta', 'Razón social correcta', 'Observación']
encabezado(ws, 4, cols, [24, 12, 12, 14, 24, 26, 50])
agr = defaultdict(lambda: [0, 0.0])
for r in combos:
    k = (r['ingenio'], r['cliente']); agr[k][0] += int(r['lineas']); agr[k][1] += float(r['cantidad'] or 0)
f = 5
for (ing, tipo), (n, cant) in sorted(agr.items(), key=lambda x: -x[1][1]):
    if tipo == 'proveedores': prop, obs = 'CEBALLOS Y LOZANO', 'Los precios de «PROVEEDOR» vienen en la tabla de Ceballos y Lozano'
    elif ing == 'pichichi': prop, obs = 'CEBALLOS Y LOZANO', 'Pichichí solo aparece en la tabla de Ceballos y Lozano'
    elif ing == 'riopaila': prop, obs = 'AGROMORALES', 'Riopaila aparece en las DOS tablas (Despeje): ¿cuándo va por Ceballos y Lozano?'
    elif ing in ('san_carlos', 'mayaguez', 'risaralda'): prop, obs = 'AGROMORALES', ''
    else: prop, obs = '', 'No aparece en ninguna tabla'
    celda(ws, f, 1, ING.get(ing, ing), color=GRIS); celda(ws, f, 2, 'Proveedor' if tipo == 'proveedores' else ('Ingenio' if tipo == 'ingenios' else '(vacío)'), color=GRIS)
    celda(ws, f, 3, n, color=GRIS); celda(ws, f, 4, cant, fmt=HA, color=GRIS); celda(ws, f, 5, prop or '—', color=GRIS)
    celda(ws, f, 6, prop or None, llenar=True); celda(ws, f, 7, obs, llenar=True)
    f += 1
lista(ws, f'F5:F{f - 1}', 'G', 3)

# ── 4 Haciendas ────────────────────────────────────────────────────────────
ws = wb.create_sheet('4 Haciendas')
titulo(ws, '4 · Haciendas: cliente y razón social de cada una',
       'Solo corrija las EXCEPCIONES. En Riopaila (Ingenio), diga si la hacienda es de Riopaila Agrícola o del Ingenio Riopaila.')
cols = ['Ingenio (en la app)', 'Tipo', 'Código hacienda', 'Nombre hacienda', 'Labores registradas', 'Hectáreas', 'Última labor',
        'Cliente que se factura', 'Razón social que factura', 'Si es proveedor: nombre o NIT a quien se factura', 'Observación']
encabezado(ws, 4, cols, [20, 11, 12, 30, 11, 11, 11, 22, 22, 30, 30])
orden = {'riopaila': 0}
filas = sorted(haciendas, key=lambda r: (orden.get(r['ingenio'], 1), r['ingenio'], r['cliente'], -float(r['ha'] or 0)))
f = 5
for r in filas:
    ing, tipo = r['ingenio'], r['cliente']
    es_prov = tipo == 'proveedores'
    if es_prov: cli, rs = 'Proveedor', 'CEBALLOS Y LOZANO'
    elif ing == 'riopaila': cli, rs = None, 'AGROMORALES'   # ← la pregunta: ¿Ingenio Riopaila o Riopaila Agrícola?
    elif ing == 'pichichi': cli, rs = 'Ingenio Pichichí', 'CEBALLOS Y LOZANO'
    elif ing in ING and ing not in ('', 'proveedor'): cli, rs = ING[ing], 'AGROMORALES'
    else: cli, rs = None, None
    celda(ws, f, 1, ING.get(ing, ing), color=GRIS); celda(ws, f, 2, 'Proveedor' if es_prov else ('Ingenio' if tipo == 'ingenios' else '(vacío)'), color=GRIS)
    celda(ws, f, 3, r['codigo_hacienda'], color=GRIS); celda(ws, f, 4, r['nombre_hacienda'], color=GRIS)
    celda(ws, f, 5, int(r['lineas']), color=GRIS); celda(ws, f, 6, float(r['ha'] or 0), fmt=HA, color=GRIS)
    celda(ws, f, 7, r['ultima'], color=GRIS)
    celda(ws, f, 8, cli, llenar=True); celda(ws, f, 9, rs, llenar=True)
    celda(ws, f, 10, None, llenar=es_prov, color=None if es_prov else GRIS); celda(ws, f, 11, None, llenar=True)
    f += 1
lista(ws, f'H5:H{f - 1}', 'B', len(CLIENTES)); lista(ws, f'I5:I{f - 1}', 'A', len(RS))
ws.auto_filter.ref = f'A4:K{f - 1}'

# ── 5 Labores y variantes ──────────────────────────────────────────────────
ws = wb.create_sheet('5 Labores y variantes')
titulo(ws, '5 · Labores y variantes', 'Cómo sabe la app qué precio le toca a cada labor. Escoja en la columna D; aclare en E si hace falta.')
cols = ['Labor en la app', 'Pregunta', 'Nuestra propuesta', 'Respuesta', 'Observación']
encabezado(ws, 4, cols, [22, 60, 55, 38, 40])
PREG = [
    ('FERTILIZACION', 'El abono tiene 3 precios (plantilla, 2x1, 4x1). ¿Cómo se sabe cuál es?', 'Crear 3 labores en la app y que el supervisor escoja al asignar'),
    ('SUBSUELO / TRIPLE', 'Para proveedores hay «normal» y «4x1». ¿Cómo se sabe cuál es?', 'Crear SUBSUELO 4X1 y TRIPLE 4X1 y que el supervisor escoja'),
    ('CULTIVO', 'Cultivo aporque: plantilla o 2x1 (proveedores). ¿Cómo se sabe?', 'Crear CULTIVO PLANTILLA y CULTIVO 2X1 y que el supervisor escoja'),
    ('REENCALLE V', '¿La «REENCALLE V» de la app es el «Reencalle caña cruda / caña verde» de la tabla?', 'Sí: se cobra con ese precio'),
    ('ACEQUIAS', '¿Se cobra por hectárea o por hectómetro? En la app hay registros de las dos formas.', 'Por hectómetro (es longitud de zanja)'),
    ('DESPEJE 0 X 0', '¿Qué es y qué precio tiene? (≈109 ha registradas en Pichichí, Riopaila, Risaralda y San Carlos)', 'Mismo precio que Despeje, salvo que digan otro'),
    ('DEVOLUCIÓN DE BASURA', 'No está en la tabla (≈226 ha en Riopaila). ¿Qué precio tiene? ¿Se factura?', 'Poner precio en la hoja 2'),
    ('OFICIOS VARIOS', '¿Cuándo es por hora ($50.000) y cuándo por jornal ($640.000)? ¿Cuántas horas tiene un jornal?', 'La app lo registra por horas; jornal = labor aparte'),
    ('Labores ya registradas', 'Las ~4.750 labores desde mayo no tienen variante. ¿Qué precio toman?', 'El de la variante más común; se puede corregir línea por línea antes de facturar'),
]
METODOS = ['El supervisor la escoge al asignar', 'Por la suerte: plantilla o soca', 'Siempre la misma', 'Sí, de acuerdo con la propuesta', 'No (ver observación)']
for i, (lab, p, prop) in enumerate(PREG, start=5):
    celda(ws, i, 1, lab, color=GRIS); celda(ws, i, 2, p, color=GRIS); celda(ws, i, 3, prop, color=GRIS)
    celda(ws, i, 4, None, llenar=True); celda(ws, i, 5, None, llenar=True)
    ws.row_dimensions[i].height = 32
lista(ws, f'D5:D{4 + len(PREG)}', 'H', len(METODOS))

# ── 6 Preguntas ────────────────────────────────────────────────────────────
ws = wb.create_sheet('6 Preguntas')
titulo(ws, '6 · Cómo se factura', 'Con esto la app arma la pre-factura igual a como la hacen hoy.')
cols = ['Pregunta', 'Por qué importa', 'Nuestra propuesta', 'Respuesta']
encabezado(ws, 4, cols, [52, 52, 42, 42])
GEN = [
    ('¿Desde qué fecha rigen estos precios?', 'Para valorar lo ya hecho desde mayo con el precio que tocaba en su fecha.', 'Desde el 16-may-2026 (primera labor en la app)'),
    ('¿Los precios incluyen IVA? ¿Qué impuestos o retenciones lleva la factura?', 'Para que el total de la pre-factura cuadre con la factura real.', 'Precio antes de IVA; la app muestra IVA aparte'),
    ('¿Cada cuánto se factura: quincenal o mensual? ¿Por cliente o por hacienda?', 'Así se agrupan las líneas en cada pre-factura.', 'Quincenal, una por razón social + cliente'),
    ('¿Se factura el área REALIZADA (la que reporta el operador) o el área de la suerte (maestro)?', 'Cambia el valor de cada línea.', 'Área realizada, como en la planilla'),
    ('¿El ingenio pide número de orden de servicio / OT / pedido en cada línea?', 'Si lo exige, la app debe pedirlo al asignar.', 'Campo opcional por labor'),
    ('¿Quién aprueba antes de facturar? ¿Se factura solo lo aprobado?', 'Evita facturar labores sin revisar.', 'Solo labores aprobadas'),
    ('¿Hay reajuste anual de tarifas? ¿En qué fecha y cómo (IPC, %)?', 'La app guarda cada precio con su vigencia; el cambio no toca lo ya facturado.', 'Se carga el precio nuevo con su fecha'),
    ('¿Quién de ustedes es el contacto para dudas de precios?', 'Para resolver rápido las líneas «sin tarifa».', ''),
    ('¿Nos pueden pasar una factura ya emitida (de cada razón social) como ejemplo?', 'Para que la pre-factura salga en el mismo formato y con los mismos datos.', 'Una de AGROMORALES y una de CEBALLOS Y LOZANO'),
]
for i, (p, por, prop) in enumerate(GEN, start=5):
    celda(ws, i, 1, p, color=GRIS); celda(ws, i, 2, por, color=GRIS); celda(ws, i, 3, prop, color=GRIS)
    celda(ws, i, 4, None, llenar=True); ws.row_dimensions[i].height = 32

# ── Listas (oculta) ────────────────────────────────────────────────────────
wl = wb.create_sheet('Listas')
for col, (nombre, valores) in enumerate([('Razón social', RS), ('Cliente', CLIENTES), ('Labor en la app', LABORES_APP),
                                         ('Unidad', UNIDADES), ('Confirma', CONFIRMA), ('Sí/No', SINO),
                                         ('Razón social (regla)', RS + ['Depende de la hacienda']), ('Método variante', METODOS)], start=1):
    wl.cell(row=1, column=col, value=nombre).font = Font(bold=True)
    for j, v in enumerate(valores, start=2):
        wl.cell(row=j, column=col, value=v)
wl.sheet_state = 'hidden'

wb.save(SALIDA)
tot = sum(resumen.values())
print(json.dumps({k: round(v, 1) for k, v in resumen.items()}, ensure_ascii=False), round(tot, 1))
print('pendientes:', len(pendientes), '· haciendas:', len(haciendas), '· salida:', SALIDA)
